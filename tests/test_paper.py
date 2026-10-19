"""论文 Word 生成测试：内容、论文结构、表格与租户目录。"""

from __future__ import annotations

from docx import Document
from docx.shared import Pt

from paper_writer import build_paper_docx, generate_paper_docx

CONTENT = """# 面向学术文献的混合检索方法研究
## 摘要
本文研究面向中文学术文献的混合检索方法，结合向量检索与 BM25，并通过交叉编码器精排。
## 关键词
混合检索；学术文献；RAG
## 1 引言
随着文献数量增长，单一向量检索难以兼顾语义与关键词匹配 [来源: 2.pdf p.4]。
## 2 方法
### 2.1 融合策略
系统分别召回候选片段，再使用加权融合排序。**实验表明**该方法有效。
| 方法 | Recall@5 | MRR |
| --- | --- | --- |
| 向量 | 0.82 | 0.71 |
| 混合 | 0.94 | 0.88 |
## 3 结论
混合检索能够提升召回与排序质量。
## 参考文献
[1] 张三. 学术文献检索研究[J]. 计算机学报, 2025.
"""


def test_build_paper_docx_structure(tmp_path):
    path = tmp_path / "paper.docx"

    build_paper_docx(
        path,
        title="面向学术文献的混合检索方法研究",
        content=CONTENT,
        author="张三",
        affiliation="示例大学 计算机学院",
    )

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "面向学术文献的混合检索方法研究" in text
    assert "摘要：" in text
    assert "关键词：" in text
    assert "1 引言" in text
    assert "2 方法" in text
    assert "参考文献" in text
    assert "[1] 张三" in text
    assert "来源:" not in text
    assert "PAGE" in document.sections[0].footer._element.xml
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 3

    heading = next(p for p in document.paragraphs if p.text == "1 引言")
    assert heading.style.name == "Heading 1"
    assert heading.runs[0].font.color.rgb is not None

    body = next(p for p in document.paragraphs if "随着文献数量增长" in p.text)
    assert body.paragraph_format.first_line_indent == Pt(24)


def test_generate_paper_docx_uses_tenant_directory(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "paper_writer.user_dir", lambda kind, user_id: tmp_path / kind / str(user_id)
    )

    path, filename = generate_paper_docx(
        7,
        title="混合检索论文",
        content=CONTENT,
        author="李四",
    )

    assert path.is_file()
    assert filename.endswith(".docx")
    assert "混合检索" in filename
    assert path.parent == tmp_path / "documents" / "7" / "papers"
