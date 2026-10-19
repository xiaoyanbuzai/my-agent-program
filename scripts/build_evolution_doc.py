"""生成《文献智库技术演进说明》DOCX。

用法：
    uv run python scripts/build_evolution_doc.py
输出：
    docs/文献智库_技术演进说明.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "文献智库_技术演进说明.docx"

ACCENT = RGBColor(0xC2, 0x41, 0x0C)
ACCENT_DARK = RGBColor(0x9A, 0x34, 0x12)
INK = RGBColor(0x2D, 0x24, 0x1B)
MUTED = RGBColor(0x8F, 0x7F, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = "FBF7F1"
LIGHT = "FDF1E7"
DARK = "2B2118"

BODY_FONT = "宋体"
HEADING_FONT = "黑体"
MONO_FONT = "Consolas"
CONTENT_WIDTH_CM = 16.2


def _set_font(
    run,
    *,
    name: str = BODY_FONT,
    size: float = 10.5,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name if name != BODY_FONT else "Times New Roman")
    rfonts.set(qn("w:hAnsi"), name if name != BODY_FONT else "Times New Roman")


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell_margins(cell, top=100, bottom=100, left=120, right=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for tag, value in (
        ("top", top),
        ("left", left),
        ("bottom", bottom),
        ("right", right),
    ):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def _bottom_border(paragraph, color: str = "E7D8C8", size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _body(
    document,
    text: str,
    *,
    indent: bool = True,
    size: float = 10.5,
    color: RGBColor = INK,
    space_after: float = 5,
    line_spacing: float = 1.35,
    bold: bool = False,
):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.first_line_indent = Pt(size * 2) if indent else Pt(0)
    run = paragraph.add_run(text)
    _set_font(run, size=size, color=color, bold=bold)
    return paragraph


def _bullet(document, text: str, *, size: float = 10.5) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.3
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(3)
    fmt.left_indent = Pt(20)
    fmt.first_line_indent = Pt(-12)
    run = paragraph.add_run("•  ")
    _set_font(run, size=size, color=ACCENT)
    run = paragraph.add_run(text)
    _set_font(run, size=size)


def _heading(document, text: str, level: int = 1) -> None:
    style = "Heading 1" if level == 1 else "Heading 2"
    paragraph = document.add_paragraph(style=style)
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(0)
    fmt.keep_with_next = True
    if level == 1:
        fmt.space_before = Pt(18)
        fmt.space_after = Pt(8)
        _set_font(
            paragraph.add_run(text),
            name=HEADING_FONT,
            size=17,
            bold=True,
            color=ACCENT_DARK,
        )
        _bottom_border(paragraph)
    else:
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(4)
        _set_font(
            paragraph.add_run(text),
            name=HEADING_FONT,
            size=12.5,
            bold=True,
            color=INK,
        )


def _small_heading(document, text: str) -> None:
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(8)
    fmt.space_after = Pt(3)
    fmt.keep_with_next = True
    _set_font(
        paragraph.add_run(text),
        name=HEADING_FONT,
        size=11,
        bold=True,
        color=INK,
    )


def _caption(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(3)
    _set_font(paragraph.add_run(text), size=9, color=MUTED, italic=False)


def _table(document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        if widths:
            cell.width = Cm(widths[index])
        _shade(cell, "9A3412")
        _cell_margins(cell, top=70, bottom=70)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_font(
            paragraph.add_run(header),
            name=HEADING_FONT,
            size=9.5,
            bold=True,
            color=WHITE,
        )

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.text = ""
            if widths:
                cell.width = Cm(widths[index])
            if row_index % 2 == 1:
                _shade(cell, PALE)
            _cell_margins(cell, top=70, bottom=70)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.2
            paragraph.paragraph_format.space_after = Pt(0)
            _set_font(paragraph.add_run(value), size=9.2)

    # 行内内容不跨页，表头会在新页重复；避免出现半行或空白行。
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _callout(document, label: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(CONTENT_WIDTH_CM)
    _shade(cell, LIGHT)
    _cell_margins(cell, top=140, bottom=140, left=180, right=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(label + "　")
    _set_font(run, name=HEADING_FONT, size=10, bold=True, color=ACCENT_DARK)
    run = paragraph.add_run(text)
    _set_font(run, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _code_block(document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(CONTENT_WIDTH_CM)
    _shade(cell, DARK)
    _cell_margins(cell, top=150, bottom=150, left=180, right=180)
    lines = text.splitlines()
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run("\n")
        _set_font(paragraph.add_run(line), name=MONO_FONT, size=8.4, color=RGBColor(0xF8, 0xEE, 0xE3))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _page_number(document) -> None:
    section = document.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    _set_font(run, size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    paragraph._p.append(field)
    run = paragraph.add_run(" 页")
    _set_font(run, size=8.5, color=MUTED)


def _header(document) -> None:
    section = document.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(
        paragraph.add_run("文献智库 · 技术演进说明"),
        size=8,
        color=MUTED,
    )
    _bottom_border(paragraph, color="E7D8C8", size=4)


def _cover(document: Document) -> None:
    for _ in range(2):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        _set_font(paragraph.add_run(""), size=20)

    line = document.add_paragraph()
    line.paragraph_format.space_after = Pt(18)
    line.paragraph_format.first_line_indent = Pt(0)
    run = line.add_run("　")
    _set_font(run, size=2)
    _bottom_border(line, color="C2410C", size=10)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    _set_font(
        paragraph.add_run("TECHNICAL EVOLUTION REPORT"),
        name=HEADING_FONT,
        size=9,
        bold=True,
        color=ACCENT,
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    _set_font(
        paragraph.add_run("文献智库"),
        name=HEADING_FONT,
        size=30,
        bold=True,
        color=INK,
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(12)
    _set_font(
        paragraph.add_run("从 RAG 原型到多租户研究工作台"),
        name=HEADING_FONT,
        size=18,
        bold=True,
        color=ACCENT_DARK,
    )

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    _set_font(
        paragraph.add_run("技术改造说明 · v0.2.0"),
        size=11,
        color=MUTED,
    )

    for _ in range(3):
        document.add_paragraph()

    details = [
        ("文档类型", "技术演进说明"),
        ("覆盖范围", "后端、检索、Agent、前端、数据模型、测试与评测"),
        ("改造原则", "保留能力、分层演进、边界先行、可验证"),
        ("更新日期", "2026 年 10 月"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    for label, value in details:
        cells = table.add_row().cells
        cells[0].width = Cm(3.5)
        cells[1].width = Cm(8.5)
        for cell in cells:
            _cell_margins(cell, top=60, bottom=60, left=0, right=80)
        _set_font(cells[0].paragraphs[0].add_run(label), size=9.5, color=MUTED, bold=True)
        _set_font(cells[1].paragraphs[0].add_run(value), size=9.5, color=INK)

    for _ in range(4):
        document.add_paragraph()

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _set_font(
        paragraph.add_run("面向学术文献的检索、对话与写作一体化实践"),
        size=10,
        color=MUTED,
        italic=True,
    )

    document.add_page_break()


def _toc(document: Document) -> None:
    _heading(document, "目录", 1)
    entries = [
        ("01", "项目概览与改造目标"),
        ("02", "改造前基线"),
        ("03", "演进路线"),
        ("04", "总体架构"),
        ("05", "关键改造"),
        ("06", "设计决策"),
        ("07", "效果与验证"),
        ("08", "模块对照"),
        ("09", "后续规划"),
    ]
    for number, title in entries:
        paragraph = document.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(4)
        fmt.space_after = Pt(4)
        fmt.first_line_indent = Pt(0)
        run = paragraph.add_run(f"{number}　")
        _set_font(run, name=HEADING_FONT, size=10.5, bold=True, color=ACCENT)
        run = paragraph.add_run(title)
        _set_font(run, size=10.5)
    document.add_paragraph()
    _callout(
        document,
        "阅读提示",
        "第 2—4 节说明背景、路线与总体结构；第 5 节按能力拆分关键改造；第 6—7 节给出设计取舍"
        "与验证数据。表格中的模块名均对应仓库中的实际文件。",
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _set_font(paragraph.add_run("关键词　"), name=HEADING_FONT, size=9, bold=True, color=ACCENT_DARK)
    _set_font(
        paragraph.add_run("多租户 · 混合检索 · 多模态 · 会话持久化 · SSE · 论文 Word"),
        size=9,
        color=MUTED,
    )
    document.add_page_break()


def build() -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)

    _header(document)
    _page_number(document)
    _cover(document)
    _toc(document)

    _heading(document, "01　项目概览与改造目标", 1)
    _body(
        document,
        "文献智库最初是一套以验证 RAG 链路为目标的学术文献助手：能够解析 PDF、建立向量索引、"
        "调用工具完成问答，并已具备多用户向量目录、图片理解与父子分块等能力。随着使用场景从"
        "“能回答问题”扩展到“多人使用、长期沉淀、可追溯引用、可交付文档”，项目需要从原型走向"
        "可部署、可维护、可验证的工程系统。",
    )
    _body(
        document,
        "本轮改造没有推倒重来，而是在保留原有能力的基础上，按“边界先行、检索提质、体验闭环、"
        "验证兜底”的顺序逐步演进。改造覆盖后端、检索、Agent、前端、数据模型、测试与评测，"
        "最终形成“检索—对话—会话沉淀—论文交付”的完整工作流。",
    )
    _callout(
        document,
        "一句话概括",
        "从单文件、以演示为目标的 RAG 原型，演进为按租户隔离、带页码引用、支持流式对话、"
        "可生成论文 Word 的研究工作台。",
    )

    _heading(document, "02　改造前基线", 1)
    _body(
        document,
        "原型阶段的核心能力已经打通：FastAPI 提供页面与接口，LangGraph 负责工具调用，"
        "Chroma 与 BM25 组成混合检索的雏形，PDF/Word/Markdown 解析、图片理解与扫描件 OCR "
        "也已接入。它证明了“文献入库—检索—回答”的可行性。",
    )
    _body(
        document,
        "当系统开始承载真实使用后，关注点从“功能是否存在”转为“边界是否可靠、结果是否可信、"
        "体验是否连续、改动是否可验证”。这构成了本轮演进的主线。",
    )
    _caption(document, "表 1　原型阶段与工程阶段的关注点对照")
    _table(
        document,
        ["维度", "原型阶段", "本轮演进方向"],
        [
            ["工程结构", "入口集中、配置分散、以单文件编排为主", "分层模块、集中配置、生命周期管理"],
            ["租户边界", "已有按用户存储的雏形", "从 JWT 到目录、检索、会话、下载全链路约束"],
            ["检索质量", "英文嵌入模型、融合与精排链路已存在", "中文嵌入、页码元数据、阈值与去重"],
            ["多模态", "图片理解与 OCR 分散在多个入口", "统一服务、类型/大小校验、模态元数据"],
            ["会话体验", "页面刷新后上下文难以恢复", "会话表、消息表、历史侧栏与切换/删除"],
            ["交付能力", "以聊天回答为主", "流式回答、论文 Word、受控下载"],
            ["验证体系", "以手工试跑为主", "单元测试、金标准评测、静态检查"],
        ],
        widths=[3.0, 6.4, 6.8],
    )

    _heading(document, "03　演进路线", 1)
    _body(
        document,
        "改造按能力依赖关系分五个阶段推进。每个阶段都保持系统可运行，并在上一阶段稳定后再进入下一层。",
    )
    _caption(document, "表 2　五个演进阶段")
    _table(
        document,
        ["阶段", "重点", "形成的成果"],
        [
            ["一、工程基础", "配置、路径、日志、生命周期、依赖管理", "config / logging / lifespan / uv 配置"],
            ["二、检索与索引", "中文嵌入、页码、父块、阈值、模型版本", "chunking / indexer / kb_retriever / vector_store"],
            ["三、多模态与文档", "图片理解、OCR、表格、文档管理", "multimodal / documents / uploads"],
            ["四、会话与界面", "会话持久化、历史侧栏、背景、响应式", "conversations / models / templates"],
            ["五、交付与交互", "SSE、论文 Word、受控下载", "SSE 路由 / paper_writer / 下载接口"],
        ],
        widths=[3.2, 6.4, 6.6],
    )

    _heading(document, "04　总体架构", 1)
    _body(
        document,
        "演进后的系统以 FastAPI 为入口，LangGraph 负责多工具编排，RAG 管线负责解析、索引与检索，"
        "SQLite 保存账号、会话与消息，Chroma 保存租户向量，knowledge 目录保存原始文件与生成文档。",
    )
    _small_heading(document, "4.1　原型阶段")
    _code_block(
        document,
        """web_app.py
  ├─ FastAPI routes
  ├─ LangGraph agent + tools
  ├─ SQLite checkpointer / store
  └─ RAG calls (loader / indexer / retriever)

单文件编排，模块边界较弱，配置与路径分散在入口中。""",
    )
    document.add_page_break()
    _small_heading(document, "4.2　工程阶段")
    _code_block(
        document,
        """Browser (JWT)
   │
   ▼
FastAPI  ── Auth / Conversations / Documents / Papers
   │
   ▼
LangGraph Agent ── kb_search / web / arXiv / paper_docx
   │
   ▼
Hybrid Retrieval ── Vector + BM25 + CrossEncoder
   │
   ▼
Per-tenant Storage
   ├─ data/chroma_db/<user_id>/
   ├─ knowledge/documents/<user_id>/
   ├─ knowledge/web/<user_id>/
   └─ data/app.db (users / conversations / messages)""",
    )
    _caption(document, "表 3　核心模块与职责")
    _table(
        document,
        ["模块", "职责"],
        [
            ["config.py", "环境变量、路径、模型、上传与检索参数、密钥策略"],
            ["uploads.py", "文件名清洗、扩展名校验、限流写盘、图片魔数"],
            ["multimodal.py", "图片理解与扫描件 OCR 的统一入口"],
            ["chunking.py", "按页/章节切块、父块截断、页码保留"],
            ["indexer.py", "父子分块、元数据、批量写入与替换语义"],
            ["vector_store.py", "租户向量库、嵌入模型版本、模型回退"],
            ["kb_retriever.py", "向量 + BM25 + 精排、阈值、父块去重"],
            ["documents.py", "文档列表、删除、原始文件清理"],
            ["conversations.py", "会话与消息持久化、标题、排序、删除"],
            ["paper_writer.py", "论文 Markdown 到中文论文格式 DOCX"],
            ["web_app.py", "路由、Agent 图、SSE、下载、生命周期"],
        ],
        widths=[4.0, 12.2],
    )

    _heading(document, "05　关键改造", 1)

    _small_heading(document, "5.1　工程基础：从分散参数到集中配置")
    _body(
        document,
        "新增 config.py，把环境变量、数据目录、模型、上传限制、检索阈值和 Agent 步数集中到一个 "
        "Settings 对象中，并通过 get_settings() 缓存读取。路径全部基于项目根目录解析，避免从不同"
        "工作目录启动时出现路径漂移。",
    )
    _bullet(document, "新增 logging_utils.py，统一控制台与轮转文件日志，替代散落的 print。")
    _bullet(document, "应用改用 lifespan 管理 SQLite 会话与记忆连接，启动和关闭路径清晰。")
    _bullet(document, "同步阻塞操作移入线程池，避免长任务阻塞事件循环。")
    _bullet(document, "新增 uv.toml，把依赖缓存指向非系统盘，并保留镜像配置。")

    _small_heading(document, "5.2　多租户：从“目录分开”到“全链路约束”")
    _body(
        document,
        "每个 user_id 拥有独立的向量库目录、原始文档目录和网页学习目录；HTTP 层统一从 JWT 的 "
        "sub 解析用户，工具内部只读取 LangGraph 的 configurable.user_id，不再把 user_id 暴露为"
        "模型可填参数。",
    )
    _bullet(document, "向量库按 user_id 物理隔离，collection 内保留 user_id 元数据用于审计。")
    _bullet(document, "文档列表与删除接口只操作当前租户的数据，删除文件前校验路径必须位于 knowledge 目录。")
    _bullet(document, "会话与消息带 user_id 归属校验，跨租户读取或删除会被拒绝。")
    _bullet(document, "论文下载接口同样按当前用户目录解析，并使用路径归一化阻止目录穿越。")

    _small_heading(document, "5.3　安全：把边界写进默认行为")
    _caption(document, "表 4　安全能力的演进")
    _table(
        document,
        ["风险面", "演进后的做法"],
        [
            ["JWT", "生产环境必须配置密钥；开发环境生成并持久化本地密钥，避免弱默认值"],
            ["上传", "扩展名白名单、分块限流写盘、大小上限、文件名清洗、图片魔数校验"],
            ["SSRF", "仅 http/https、禁止 URL 凭证、限制端口、DNS 后拒绝非公网地址、重定向二次校验"],
            ["XSS", "模型与文档内容统一经 DOMPurify；缺失时退回纯文本转义"],
            ["提示注入", "system prompt 明确检索内容只是资料，不是指令"],
            ["异常信息", "前端只收到通用提示，堆栈写入服务端日志"],
            ["下载", "论文文件必须携带 JWT，前端以 fetch + blob 方式取回"],
        ],
        widths=[3.2, 13.0],
    )

    _small_heading(document, "5.4　多模态：从分散调用到统一服务")
    _body(
        document,
        "图片理解与扫描件 OCR 归入 multimodal.py，统一处理输入校验、超时与异常。图片支持 PNG / "
        "JPEG / WebP，识别结果以 modality=image 入库，可被文本问题检索；扫描版 PDF 按页 OCR，"
        "单页异常不阻断整篇；Word 表格转为 Markdown 入库，保留结构化信息。",
    )

    _small_heading(document, "5.5　检索：从“能召回”到“可解释、可评估”")
    _caption(document, "表 5　检索链路演进")
    _table(
        document,
        ["环节", "演进后的设计"],
        [
            ["嵌入", "BAAI/bge-base-zh-v1.5，查询侧加 BGE 指令前缀，文档侧不加"],
            ["切块", "子块 400 字用于召回，父块按句号截断到 1500 字用于生成"],
            ["元数据", "source / page / page_end / modality / parent_id / file_hash"],
            ["融合", "向量与 jieba BM25 加权融合，再由 CrossEncoder 精排"],
            ["阈值", "精排分数可配置阈值，低相关结果不进入上下文"],
            ["去重", "按 parent_id 去重，避免同一父块重复占用上下文"],
            ["版本", "collection 记录 embedding_model，换模型后提示重建索引"],
            ["写入", "先写新块，成功后再删旧块；批量写入并保留安全恢复路径"],
        ],
        widths=[3.0, 13.2],
    )
    _body(
        document,
        "页码是引用可信度的基础：无书签 PDF 逐页切块；有书签 PDF 按一级章节切并记录起始页。"
        "回答中的引用可以落到“文件名 + 页码”，而不是只给一个模糊的文档名。",
    )

    _small_heading(document, "5.6　Agent：多工具协作与边界控制")
    _body(
        document,
        "Agent 工具集扩展为知识库检索、arXiv 检索、网页学习、联网搜索和论文 Word 生成。"
        "system prompt 统一约束检索纪律、引用格式和拒答条件；长期记忆按问题相关性排序后注入，"
        "避免无关历史污染上下文。",
    )
    _bullet(document, "单轮步数预算提升到 24，满足“检索 + 生成 + 回答”的多工具任务。")
    _bullet(document, "论文生成模式限制检索轮次，检索后进入文档生成，工具成功后立即收束回答。")
    _bullet(document, "工具调用记录名称与参数摘要，便于审计与性能分析。")

    _small_heading(document, "5.7　会话与界面：从一次性页面到研究工作台")
    _body(
        document,
        "新增 conversations 与 chat_messages 表，会话和消息落入 SQLite。页面加载时拉取会话列表，"
        "按 updated_at 倒序展示；选择会话后从数据库恢复消息；刷新或重新登录不会丢失历史。",
    )
    _bullet(document, "左侧历史侧栏支持新建、切换、删除会话；首条用户问题自动成为标题。")
    _bullet(document, "删除会话时同步清理 LangGraph checkpoint，避免残留线程。")
    _bullet(document, "五套内置暖色背景与本地图片壁纸；上传图片自动压缩后存入 localStorage。")
    _bullet(document, "响应式布局：桌面为侧栏 + 主对话区，移动端为抽屉式侧栏。")
    _bullet(document, "文档管理改为弹窗，支持查看与删除当前账号的入库内容。")

    _small_heading(document, "5.8　论文 Word：从聊天回答到可交付文档")
    _body(
        document,
        "新增 paper_writer.py，把论文 Markdown 排版为 .docx。Agent 在识别到“生成论文 doc / 导出 "
        "Word”等意图后，先检索知识库，再调用 generate_paper_docx_tool，生成文件保存到当前租户的 "
        "papers 目录，并返回受控下载链接。",
    )
    _caption(document, "表 6　论文 DOCX 的版式约定")
    _table(
        document,
        ["元素", "格式"],
        [
            ["页面", "A4，上下左右 2.54 cm"],
            ["标题", "黑体二号，居中"],
            ["作者/单位", "宋体居中，单位使用灰色小字"],
            ["摘要/关键词", "加粗标签 + 正文，摘要 200–300 字"],
            ["章节标题", "一级黑体四号，二级黑体小四"],
            ["正文", "宋体小四，1.5 倍行距，首行缩进 2 字符，两端对齐"],
            ["表格", "Table Grid，表头加粗居中"],
            ["参考文献", "宋体五号，悬挂缩进"],
            ["页码", "页脚居中，自动 PAGE 域"],
        ],
        widths=[3.4, 12.8],
    )
    _body(
        document,
        "生成时会清理正文中的“[来源: …]”检索标记，论文内引用改用 [1]、[2] 并在文末参考文献列出，"
        "使文档更接近正式论文，而不是检索日志。",
    )

    _small_heading(document, "5.9　流式输出：让长回答可见地生成")
    _body(
        document,
        "新增 POST /ask/stream，以 SSE 返回 token、done 与 error 事件；前端使用 fetch + "
        "ReadableStream 解析，按 80ms 节流渲染 Markdown，结束后统一触发 MathJax。非流式 /ask "
        "保留，便于脚本调用和兼容。",
    )

    _small_heading(document, "5.10　测试与评测：把“感觉可用”变成“可以度量”")
    _body(
        document,
        "新增 37 个 pytest 测试，覆盖上传安全、JWT、SSRF、切块页码、检索阈值与去重、索引元数据、"
        "租户隔离、文档删除、会话持久化、论文排版与 SSE 生成器。",
    )
    _body(
        document,
        "新增 eval/golden_set.json 与 run_eval.py，用金标准问题集计算 hit@k、MRR、页码命中率和"
        "平均延迟，为后续模型替换提供回归基线。",
    )

    _heading(document, "06　设计决策", 1)
    _caption(document, "表 7　关键设计取舍")
    _table(
        document,
        ["决策", "原因"],
        [
            ["每个租户独立 Chroma 目录", "比单 collection + where 过滤更直接地避免跨租户检索"],
            ["父子块 + 父块截断", "兼顾召回粒度与上下文成本，避免整章进入提示词"],
            ["向量 + BM25 + 精排", "语义召回与关键词匹配互补，精排负责最终相关性"],
            ["页码写入元数据", "引用可回溯到具体页，支持审阅和核验"],
            ["会话与 checkpoint 分离", "业务消息用于展示，checkpoint 用于模型上下文"],
            ["同步生成器承载 SSE", "兼容现有 SQLite checkpointer，同时不阻塞事件循环"],
            ["壁纸存 localStorage", "图片只属于浏览器偏好，不增加服务器存储与隐私面"],
        ],
        widths=[5.0, 11.2],
    )

    _heading(document, "07　效果与验证", 1)
    _body(
        document,
        "在当前测试语料与本地环境下，系统完成了从检索到交付的闭环验证。以下数据用于说明当前基线，"
        "不代表所有语料的泛化表现。",
    )
    _caption(document, "表 8　当前验证结果")
    _table(
        document,
        ["指标", "结果", "说明"],
        [
            ["知识库规模", "306 个子块", "用户 1 的两篇 PDF"],
            ["hit@4", "1.00", "6 题内部金标准集"],
            ["MRR", "1.00", "首个相关结果排名"],
            ["page_hit", "1.00", "top-k 内命中期望页码"],
            ["平均检索延迟", "约 2.95 s", "本地 CPU，含精排"],
            ["测试", "37 passed", "pytest"],
            ["静态检查", "All checks passed", "ruff"],
            ["流式输出", "1463 个 token 事件 + done", "一次真实问答采样"],
            ["论文生成", "8 页 DOCX", "含标题、摘要、章节、参考文献、页码"],
        ],
        widths=[4.2, 5.4, 6.6],
    )
    _body(
        document,
        "这组结果说明：检索链路已经具备可量化的回归基线，会话与文档能力已经形成可演示的闭环，"
        "后续替换精排模型或扩充语料时，可以直接对比同一套指标。",
    )

    _heading(document, "08　模块对照", 1)
    _caption(document, "表 9　主要文件的前后对照")
    _table(
        document,
        ["原型阶段", "工程阶段", "变化"],
        [
            ["web_app.py（单文件）", "web_app.py + 多个领域模块", "职责拆分，入口保留编排"],
            ["配置散落在入口", "config.py", "集中配置与路径"],
            ["解析/切块混合", "kb_loader.py + chunking.py", "解析与切块分离"],
            ["索引逻辑内嵌", "indexer.py + vector_store.py", "写入与存储解耦"],
            ["检索逻辑内嵌", "kb_retriever.py", "混合检索与阈值独立"],
            ["无文档管理", "documents.py", "列表、删除、文件清理"],
            ["无会话表", "conversations.py + models.py", "会话与消息持久化"],
            ["无论文导出", "paper_writer.py", "Markdown 到 DOCX"],
            ["页面脚本为主", "templates/index.html + login.html", "工作台、历史、壁纸、SSE"],
            ["手工试跑", "tests/ + eval/", "自动测试与评测"],
        ],
        widths=[4.8, 6.2, 5.2],
    )

    _heading(document, "09　后续规划", 1)
    _body(
        document,
        "当前版本已经形成稳定的能力底座。下一阶段重点是把“可用”继续推进到“更好用、更可信、更易维护”。",
    )
    _bullet(document, "精排模型切换到中文 CrossEncoder，进一步改善中文相关性判断。")
    _bullet(document, "BM25 改为增量索引，降低文档增长后的重建成本。")
    _bullet(document, "为流式回答增加“正在检索 / 正在阅读第 X 页”的工具状态提示。")
    _bullet(document, "论文模板支持期刊/学位论文切换，并支持分章节生成后合并。")
    _bullet(document, "扩充金标准集与跨语料评测，形成持续回归的质量看板。")

    document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(
        paragraph.add_run("— 让文献检索、对话与写作发生在同一个工作台 —"),
        size=9.5,
        color=MUTED,
        italic=True,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
