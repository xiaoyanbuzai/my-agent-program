"""多格式文档解析:PDF / Word / Markdown / txt → 带页码与模态的结构化块。"""

from __future__ import annotations

import logging
import re
from pathlib import Path

import pymupdf

from chunking import split_by_chars, split_by_toc, split_text_blocks
from config import get_settings
from multimodal import MultimodalError, ocr_pdf_page

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}
_ZERO_WIDTH = re.compile(r"[\u200b-\u200f\u2028\u2029\ufeff]")
_PAGE_NUMBER_LINE = re.compile(r"^\s*[-–—]?\s*\d{1,4}\s*[-–—]?\s*$", re.MULTILINE)


class DocumentParseError(ValueError):
    """文档无法解析。"""


def load_document(file_path: str) -> list[dict]:
    """解析文档，返回 ``[{"title", "content", "page", "modality"}, ...]``。"""
    path = Path(file_path)
    if not path.is_file():
        raise DocumentParseError(f"文件不存在：{path.name}")

    ext = path.suffix.lower()
    if ext not in TEXT_EXTENSIONS:
        raise DocumentParseError(f"不支持的文件类型：{ext}")

    if ext == ".pdf":
        blocks = _load_pdf(path)
    elif ext == ".docx":
        blocks = _load_docx(path)
    elif ext in (".md", ".markdown"):
        blocks = _load_markdown(path)
    else:
        blocks = _load_txt(path)

    cleaned: list[dict] = []
    for block in blocks:
        content = clean_raw_text(block.get("content", ""))
        if not content:
            continue
        cleaned.append(
            {
                "title": (block.get("title") or "正文").strip() or "正文",
                "content": content,
                "page": int(block.get("page") or 0),
                "modality": block.get("modality") or "text",
            }
        )
    return cleaned


def _load_pdf(path: Path) -> list[dict]:
    """PDF:优先文本层；文本过少时判定为扫描件并调用视觉模型 OCR。"""
    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise DocumentParseError("PDF 打开失败或文件已损坏") from exc

    try:
        full_text = "\n".join((page.get_text() or "") for page in doc)
        if len(full_text.strip()) < 50 and doc.page_count > 0:
            return _ocr_pdf(doc, path)

        toc = doc.get_toc()
        return split_by_toc(doc, toc) if toc else split_by_chars(doc)
    finally:
        doc.close()


def _ocr_pdf(doc: pymupdf.Document, path: Path) -> list[dict]:
    settings = get_settings()
    max_pages = min(doc.page_count, settings.ocr_max_pages)
    logger.info("检测到扫描版 PDF %s,OCR 前 %s 页", path.name, max_pages)

    blocks: list[dict] = []
    succeeded = 0
    for page_index in range(max_pages):
        try:
            pixmap = doc[page_index].get_pixmap(dpi=150)
            text = ocr_pdf_page(pixmap.tobytes("png"))
            succeeded += 1
        except MultimodalError as exc:
            logger.warning("第 %s 页 OCR 失败(%s):%s", page_index + 1, exc, path.name)
            text = "[该页识别失败]"
        except Exception:
            logger.exception("第 %s 页 OCR 失败：%s", page_index + 1, path.name)
            text = "[该页识别失败]"
        blocks.append(
            {
                "title": f"第 {page_index + 1} 页",
                "content": text,
                "page": page_index + 1,
                "modality": "text",
            }
        )

    if succeeded == 0:
        raise DocumentParseError(
            "扫描版 PDF 需要视觉模型 OCR,请检查 ZHIPU_API_KEY 配置或稍后重试"
        )

    if doc.page_count > max_pages:
        blocks.append(
            {
                "title": "OCR 范围提示",
                "content": f"该 PDF 共 {doc.page_count} 页，仅自动识别了前 {max_pages} 页。",
                "page": max_pages,
                "modality": "text",
            }
        )
    return blocks


def _load_docx(path: Path) -> list[dict]:
    """Word:按 Heading 样式分节，并把表格转成 Markdown 保留数据。"""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("缺少 python-docx 依赖，无法解析 docx") from exc

    doc = Document(str(path))
    sections: list[tuple[str, list[str]]] = []
    current_title = "正文"
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_title, current_lines
        text = "\n".join(line for line in current_lines if line.strip())
        if text.strip():
            sections.append((current_title, [text]))
        current_lines = []

    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = paragraph.style
        style_name = ((style.name if style else "") or "").lower()
        style_id = (getattr(style, "style_id", "") or "") if style else ""
        is_heading = "heading" in style_name or style_id in {
            "Heading1",
            "Heading2",
            "Heading3",
            "Title",
        }
        if is_heading:
            flush()
            current_title = text
        else:
            current_lines.append(text)
    flush()

    if not sections:
        full = "\n".join(p.text for p in doc.paragraphs if (p.text or "").strip())
        if full.strip():
            sections.append(("全文", [full]))

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append("| " + " | ".join(cells) + " |")
        if len(rows) >= 2:
            header, body = rows[0], rows[1:]
            markdown = "\n".join([header, "|" + " --- |" * len(table.columns), *body])
            sections.append((f"表格 {table_index}", [markdown]))

    blocks: list[dict] = []
    for title, lines in sections:
        blocks.extend(split_text_blocks("\n".join(lines), title=title))
    return blocks


def _load_markdown(path: Path) -> list[dict]:
    """Markdown:按标题切节，保留正文；代码块围栏不进入正文。"""
    text = _read_text(path)
    sections: list[tuple[str, list[str]]] = []
    current_title = "文档开头"
    current_lines: list[str] = []
    in_code = False

    def flush() -> None:
        nonlocal current_title, current_lines
        body = "".join(current_lines).strip()
        if body:
            sections.append((current_title, [body]))
        current_lines = []

    for raw in text.splitlines(keepends=True):
        stripped = raw.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if not in_code and stripped.startswith("#"):
            flush()
            current_title = stripped.lstrip("#").strip() or "未命名小节"
            continue
        current_lines.append(raw)
    flush()

    blocks: list[dict] = []
    if sections:
        for title, lines in sections:
            blocks.extend(split_text_blocks("".join(lines), title=title))
    else:
        blocks.extend(split_text_blocks(text, title="全文"))
    return blocks


def _load_txt(path: Path) -> list[dict]:
    return split_text_blocks(_read_text(path), title="全文")


def _read_text(path: Path) -> str:
    """中文文本常见编码兜底：utf-8-sig → utf-8 → gb18030。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def clean_raw_text(text: str) -> str:
    """通用清洗：去零宽字符、断行连字符、页码行、多余空行。"""
    if not text:
        return ""
    text = _ZERO_WIDTH.sub("", text)
    text = re.sub(r"(\w)-\n(?=[a-zA-Z])", r"\1", text)  # 英文断行连字符
    text = _PAGE_NUMBER_LINE.sub("", text)
    text = re.sub(r"[-=]{6,}", "", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    # 【新增拦截】：一旦遇到单独成行的 References 或 参考文献，直接砍掉后面的所有内容
    ref_pattern = re.compile(
        r"\n(?:References|Bibliography|REFERENCES|参考文献|引用文献)\s*\n", 
        re.IGNORECASE
    )
    match = ref_pattern.search(text)
    if match:
        text = text[:match.start()]
        
    return text.strip()


if __name__ == "__main__":
    import sys

    for argument in sys.argv[1:]:
        print(f"== {argument} ==")
        try:
            parsed = load_document(argument)
        except DocumentParseError as error:
            print(f"  解析失败：{error}")
            continue
        print(f"  blocks: {len(parsed)}")
        for block in parsed[:3]:
            print(f"  - [p{block['page']}] {block['title']} ({len(block['content'])} 字)")
